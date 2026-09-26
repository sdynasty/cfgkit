#!/usr/bin/env python3
"""生成示例 Excel 配置模板。

产物（excel/ 目录下）:
  common.xlsx  - EnumItemQuality 枚举定义表
  item.xlsx    - Item 道具表
  monster.xlsx - Monster 怪物表

表头规范（数据表 sheet，固定 4 行表头，第 5 行起为数据）:
  行1 字段名(英文)  行2 类型  行3 中文注释  行4 导出标记(cs/c/s/-)
枚举表 sheet 以 "Enum" 开头，表头 1 行: name / value / comment
数据行首列以 '#' 开头的整行视为注释行，导表跳过。
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "excel")
os.makedirs(OUT, exist_ok=True)

HEADER_STYLE = [
    # (背景色, 字体)
    (PatternFill("solid", fgColor="D9D9D9"), Font(bold=True)),                    # 行1 字段名
    (PatternFill("solid", fgColor="BDD7EE"), Font(name="Menlo", size=10)),        # 行2 类型
    (PatternFill("solid", fgColor="C6EFCE"), Font(size=10)),                      # 行3 注释
    (PatternFill("solid", fgColor="FFE699"), Font(size=10)),                      # 行4 导出标记
]
THIN = Border(*[Side(style="thin", color="AAAAAA")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")
NOTE_FONT = Font(color="999999", italic=True)


def write_table_sheet(ws, fields, rows, widths=None):
    """fields: [(字段名, 类型, 注释, 导出标记)]; rows: 数据行(list) 或 '#' 开头的注释行(str)"""
    for c, (name, typ, comment, flag) in enumerate(fields, 1):
        for r, (val, (fill, font)) in enumerate(
            zip((name, typ, comment, flag), HEADER_STYLE), 1
        ):
            cell = ws.cell(r, c, val)
            cell.fill = fill
            cell.font = font
            cell.border = THIN
            cell.alignment = CENTER
    for r, row in enumerate(rows, 5):
        if isinstance(row, str):  # 注释行
            cell = ws.cell(r, 1, row)
            cell.font = NOTE_FONT
            continue
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.border = THIN
    for c in range(1, len(fields) + 1):
        w = (widths or {}).get(c - 1, 16)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A5"


def write_enum_sheet(ws, entries, title="品质"):
    """entries: [(name, value, comment)]"""
    for c, h in enumerate(("name", "value", "comment"), 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font = HEADER_STYLE[0]
        cell.border = THIN
        cell.alignment = CENTER
    for r, (name, value, comment) in enumerate(entries, 2):
        for c, val in enumerate((name, value, comment), 1):
            cell = ws.cell(r, c, val)
            cell.border = THIN
    for c, w in enumerate((16, 10, 24), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


def write_struct_sheet(ws, fields):
    """fields: [(name, type, comment)]，内联结构体定义，每行一个字段"""
    for c, h in enumerate(("name", "type", "comment"), 1):
        cell = ws.cell(1, c, h)
        cell.fill, cell.font = HEADER_STYLE[0]
        cell.border = THIN
        cell.alignment = CENTER
    for r, (name, typ, comment) in enumerate(fields, 2):
        for c, val in enumerate((name, typ, comment), 1):
            cell = ws.cell(r, c, val)
            cell.border = THIN
            if c == 2:
                cell.font = Font(name="Menlo", size=10)
    for c, w in enumerate((14, 18, 30), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------- common.xlsx
wb = Workbook()
ws = wb.active
ws.title = "EnumItemQuality"
write_enum_sheet(ws, [
    ("White",  1, "普通(白)"),
    ("Green",  2, "优秀(绿)"),
    ("Blue",   3, "精良(蓝)"),
    ("Purple", 4, "史诗(紫)"),
    ("Orange", 5, "传说(橙)"),
])
ws = wb.create_sheet("StructDrop")
write_struct_sheet(ws, [
    ("item",  "ref<Item>", "掉落道具ID"),
    ("rate",  "float",     "掉落概率(0~1)"),
    ("count", "int",       "掉落数量"),
    ("weight", "int?=1",   "权重(可选,省略时取默认值1)"),
])
ws = wb.create_sheet("StructPos")
write_struct_sheet(ws, [
    ("x", "int", "X坐标"),
    ("y", "int", "Y坐标"),
])
wb.save(os.path.join(OUT, "common.xlsx"))

# ------------------------------------------------------------------ item.xlsx
wb = Workbook()
ws = wb.active
ws.title = "Item"
write_table_sheet(
    ws,
    fields=[
        ("id",        "int",               "道具ID(主键)",                      "cs"),
        ("name",      "string#uniq",       "道具名(唯一索引)",                  "cs"),
        ("quality",   "enum<ItemQuality>", "品质,见common.xlsx/EnumItemQuality", "cs"),
        ("price",     "int",               "售价(铜币)",                        "cs"),
        ("stack_max", "int",               "堆叠上限",                          "s"),
        ("desc",      "string",            "道具描述",                          "c"),
    ],
    rows=[
        "# ===== 武器 =====",
        [1001, "木剑",   "White",  100,   99,  "新手练习用的木剑"],
        [1002, "铁剑",   "Green",  500,   99,  "锋利的铁制长剑"],
        [1003, "火焰剑", "Blue",   2000,  99,  "附带火焰的长剑"],
        [1004, "屠龙刀", "Purple", 10000, 10,  "传说中的神兵"],
        "# ===== 药水 =====",
        [1005, "治疗药水", "White", 50,   999, "恢复100点生命"],
        [1006, "法力药水", "Green", 80,   999, "恢复50点法力"],
    ],
    widths={0: 10, 1: 14, 2: 20, 3: 10, 4: 10, 5: 34},
)
wb.save(os.path.join(OUT, "item.xlsx"))

# --------------------------------------------------------------- monster.xlsx
wb = Workbook()
ws = wb.active
ws.title = "Monster"
write_table_sheet(
    ws,
    fields=[
        ("id",         "int",                  "怪物ID(主键)",               "cs"),
        ("name",       "string#index",         "怪物名(普通索引,可重名)",    "cs"),
        ("level",      "int",                  "等级",                       "cs"),
        ("hp",         "int",                  "生命值",                     "cs"),
        ("drops",      "list<ref<Item>>",      "掉落道具ID,用|分隔",         "s"),
        ("tags",       "list<string>",         "标签,用|分隔",               "cs"),
        ("ai_params",  "map<string,int>",      "AI参数,如 aggro:1;speed:3",  "s"),
        ("born_pos",   "struct<Pos>",          "出生坐标,如 x:100;y:200",    "cs"),
        ("drop_group", "list<struct<Drop>>",   "权重掉落组,整组用|分隔",     "s"),
        ("desc",       "string",               "描述",                       "c"),
    ],
    rows=[
        [2001, "野狼",     3,  150,  "1001|1005",      "野兽",      "aggro:1;speed:3", "x:120;y:340",  "item:1001;rate:0.6;count:1|item:1005;rate:0.3;count:2", "新手村外游荡的野狼"],
        [2002, "哥布林",   5,  300,  "1002|1005|1006", "人形",      "aggro:2;speed:2", "x:88;y:210",   "item:1002;rate:0.4;count:1|item:1006;rate:0.4;count:1", "弱小但成群出没"],
        [2003, "火焰蜥蜴", 8,  650,  "1003",           "野兽|火系", "aggro:3;speed:1", "x:450;y:920",  "item:1003;rate:0.25;count:1",                           "尾部可喷出火焰"],
        [2004, "幼龙",     15, 3000, "1004|1003",      "龙类|BOSS", "aggro:5;speed:4", "x:1024;y:777", "item:1004;rate:0.05;count:1;weight:3|item:1003;rate:0.5;count:1", "龙穴的守护者"],
        [2005, "宝箱怪",   10, 1200, "1002|1006",      "拟态",      "stealth:5;aggro:4", "x:66;y:130", "item:1002;rate:0.8;count:1|item:1006;rate:0.8;count:2", "伪装成宝箱的魔物"],
    ],
    widths={0: 10, 1: 14, 2: 8, 3: 10, 4: 20, 5: 14, 6: 22, 7: 16, 8: 46, 9: 30},
)
wb.save(os.path.join(OUT, "monster.xlsx"))

print("模板已生成 ->", OUT)
for f in sorted(os.listdir(OUT)):
    print("  -", f)
